"""Shoonya phone bot - trading engine (p4a safe, no pandas)."""
import os
import time
import threading
import yaml
from datetime import datetime
from openpyxl import load_workbook
from NorenRestApiPy.NorenApi import NorenApi


class ShoonyaApiPy(NorenApi):
    def __init__(self):
        super().__init__(
            host="https://trade.shoonya.com/NorenWClientWeb/",
            websocket="wss://trade.shoonya.com/NorenWSWeb/"
        )


class TradingBot:
    def __init__(self, on_log=None, on_pnl=None, on_positions=None,
                 on_status=None, on_order=None):
        self.api = None
        self.token_df = []
        self.super_token = ""
        self.trading_cap = 1000
        self.userId = ""
        self.running = False
        self._thread = None
        self._stop_event = threading.Event()
        self.processed_stocks = set()

        self.on_log = on_log or (lambda *a, **k: None)
        self.on_pnl = on_pnl or (lambda *a, **k: None)
        self.on_positions = on_positions or (lambda *a, **k: None)
        self.on_status = on_status or (lambda *a, **k: None)
        self.on_order = on_order or (lambda *a, **k: None)

    def _log(self, msg):
        ts = datetime.now().strftime("%H:%M:%S")
        self.on_log(f"[{ts}] {msg}")

    def load_excel(self, path):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = self.token_df = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or r[4] is None:
                continue
            sym = str(r[4]).strip().upper()
            if sym:
                rows.append({"TradingSymbol": sym, "Token": str(r[1]).strip()})
        wb.close()
        self._log(f"Loaded {len(rows)} symbols from Excel")

    def login(self, super_token, trading_cap, password=None, user_id=None):
        from . import config as cfg
        user_id_v, pwd_v = user_id, password
        if (not user_id_v or not pwd_v) and os.path.exists(cfg.BROKER_CRED):
            with open(cfg.BROKER_CRED) as f:
                cred = yaml.load(f, Loader=yaml.FullLoader)
            user_id_v = user_id_v or cred.get("user")
            pwd_v = pwd_v or cred.get("pwd")

        if not user_id_v or not pwd_v:
            self._log("Missing userId/password - provide in dashboard")
            self.on_status("LOGIN_FAILED")
            return False

        self.userId = user_id_v
        self.super_token = super_token
        self.trading_cap = trading_cap

        self._log("Logging into Shoonya broker...")
        self.api = ShoonyaApiPy()
        ret = self.api.set_session(self.userId, pwd_v, super_token)

        if ret:
            if not os.path.exists(cfg.BROKER_CRED) and user_id and password:
                try:
                    os.makedirs(os.path.dirname(cfg.BROKER_CRED), exist_ok=True)
                    with open(cfg.BROKER_CRED, "w") as f:
                        yaml.dump({"user": user_id, "pwd": password}, f)
                except Exception:
                    pass
            self._log(f"Login successful | User: {self.userId} | Capital: {self.trading_cap}")
            self.on_status("CONNECTED")
            return True
        else:
            self._log("Login FAILED")
            self.on_status("LOGIN_FAILED")
            return False

    def start(self, super_token, trading_cap, password=None, user_id=None):
        if self.running:
            self._log("Bot already running")
            return False

        if not self.load_excel_safe():
            return False

        if not self.login(super_token, trading_cap, password, user_id):
            return False

        self._stop_event.clear()
        self.running = True
        self.processed_stocks = set()
        self._thread = threading.Thread(target=self._monitor_loop, daemon=True)
        self._thread.start()
        self.on_status("RUNNING")
        self._log("Monitor thread started")
        return True

    def load_excel_safe(self):
        from . import config as cfg
        if not os.path.exists(cfg.EXCEL_FILE):
            self._log(f"ERROR: Excel not found at {cfg.EXCEL_FILE}")
            return False
        try:
            self.load_excel(cfg.EXCEL_FILE)
            return True
        except Exception as e:
            self._log(f"ERROR loading Excel: {e}")
            return False

    def stop(self):
        if not self.running:
            return
        self._stop_event.set()
        self.running = False
        self.on_status("STOPPED")
        self._log("Bot stopped")

    def is_time_in_valid_window(self, time_str):
        from .config import VALID_TIME_WINDOWS
        check_time = datetime.strptime(time_str, "%H:%M").time()
        for start_str, end_str in VALID_TIME_WINDOWS:
            start_time = datetime.strptime(start_str, "%H:%M").time()
            end_time = datetime.strptime(end_str, "%H:%M").time()
            if start_time <= check_time <= end_time:
                return True
        return False

    def parse_webhook_payload(self, data):
        if not data:
            raise ValueError("Webhook payload is empty")

        stocks_raw = data.get("stocks", "")
        prices_raw = data.get("trigger_prices", "")
        triggered_at = data.get("triggered_at", "")

        if not stocks_raw or not prices_raw or not triggered_at:
            raise ValueError("Missing required fields")

        symbols = [s.strip() for s in stocks_raw.split(",") if s.strip()]
        prices = [p.strip() for p in prices_raw.split(",") if p.strip()]

        if len(symbols) != len(prices):
            raise ValueError("Stock and price count mismatch")

        from .config import MAX_STOCKS_PER_TIME
        if len(symbols) > MAX_STOCKS_PER_TIME:
            self._log(f"Received {len(symbols)} stocks. Max is {MAX_STOCKS_PER_TIME}. Skipping.")
            return []

        trigger_time = datetime.strptime(
            triggered_at.strip().upper(), "%I:%M %p"
        ).strftime("%H:%M")

        parsed = []
        for symbol, price in zip(symbols, prices):
            parsed.append({
                "symbol": symbol,
                "trigger_price": float(price),
                "time": trigger_time,
            })
        return parsed

    def filter_stocks(self, stocks):
        if not stocks:
            return []
        for stock in stocks:
            time_value = stock.get("time", "").strip()
            if not self.is_time_in_valid_window(time_value):
                self._log(f"Time {time_value} not in valid window")
                return []
        return stocks

    def place_sell_orders(self, stocks):
        self._log("--- PLACING SELL ORDERS ---")
        for stock in stocks:
            symbol = stock.get("symbol", "UNKNOWN").strip()
            symbol_eq = f"{symbol}-EQ"

            try:
                matches = [m for m in self.token_df if m["TradingSymbol"] == symbol_eq.upper()]
                if not matches:
                    self._log(f"FAILED {symbol_eq}: Token not found")
                    continue

                token = matches[0]["Token"]
                quote = self.api.get_quotes(exchange="NSE", token=token)

                if not quote or quote.get("stat") != "Ok":
                    self._log(f"FAILED {symbol_eq}: Quote failed")
                    continue

                current_ltp = float(quote.get("lp", 0))
                if current_ltp <= 0:
                    self._log(f"FAILED {symbol_eq}: Invalid price")
                    continue

                quantity = int(self.trading_cap / current_ltp)

                response = self.api.place_order(
                    buy_or_sell="S",
                    product_type="I",
                    exchange="NSE",
                    tradingsymbol=symbol_eq,
                    quantity=quantity,
                    discloseqty=0,
                    price_type="MKT",
                    retention="DAY",
                    remarks="Webhook_Sell",
                )

                order_id = response.get("norenordno") if response else None
                if order_id:
                    msg = f"SELL | {symbol_eq} | Price: {current_ltp:.2f} | Qty: {quantity} | Order: {order_id}"
                    self._log(msg)
                    self.on_order({"symbol": symbol_eq, "direction": "SELL", "price": current_ltp, "qty": quantity, "order_id": order_id})
                else:
                    self._log(f"FAILED {symbol_eq}: {response}")

            except Exception as e:
                self._log(f"ERROR {symbol}: {e}")

    def handle_webhook(self, data):
        self._log("=" * 60)
        self._log(f"WEBHOOK RECEIVED - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        self._log("=" * 60)

        try:
            parsed = self.parse_webhook_payload(data)
            filtered = self.filter_stocks(parsed)
            if not filtered:
                self._log("No stocks passed filter")
                return {"status": "no_match"}
            self.place_sell_orders(filtered)
            return {"status": "success"}
        except Exception as e:
            self._log(f"Webhook error: {e}")
            return {"status": "error", "message": str(e)}

    def get_total_pnl(self):
        try:
            ret = self.api.get_positions()
            if not ret:
                return 0.0
            mtm = sum(float(i.get("urmtom", 0)) for i in ret)
            pnl = sum(float(i.get("rpnl", 0)) for i in ret)
            return round(mtm + pnl, 2)
        except Exception as e:
            self._log(f"PnL Fetch Error: {e}")
            return 0.0

    def get_positions(self):
        try:
            return self.api.get_positions() or []
        except Exception:
            return []

    def _monitor_loop(self):
        from .config import PNL_STOP_LOSS_PCT, PNL_TAKE_PROFIT_PCT, MONITOR_INTERVAL
        cycle = 1
        while not self._stop_event.is_set():
            total_pnl = self.get_total_pnl()
            self.on_pnl(total_pnl)

            try:
                positions = self.api.get_positions()
                if not positions:
                    self.on_positions([])
                    self._stop_event.wait(MONITOR_INTERVAL)
                    cycle += 1
                    continue

                pos_list = []
                for p in positions:
                    sym = str(p.get("tsym", "")).strip().upper()
                    qty = int(float(p.get("netqty", 0)))
                    urmtom = float(p.get("urmtom", 0))
                    pos_list.append({"symbol": sym, "qty": qty, "pnl": urmtom})
                    self._log(f"{sym} | Qty: {qty} | PnL: {urmtom}")

                    if sym in self.processed_stocks or qty == 0:
                        continue

                    lower = -(self.trading_cap * PNL_STOP_LOSS_PCT)
                    upper = self.trading_cap * PNL_TAKE_PROFIT_PCT

                    if urmtom <= lower or urmtom >= upper:
                        try:
                            response = self.api.place_order(
                                buy_or_sell="B",
                                product_type="I",
                                exchange="NSE",
                                tradingsymbol=sym,
                                quantity=abs(qty),
                                discloseqty=0,
                                price_type="MKT",
                                retention="DAY",
                                remarks="Exit_Order",
                            )
                            if response and response.get("norenordno"):
                                self.processed_stocks.add(sym)
                                self._log(f"EXITED {sym}")
                                self.on_order({"symbol": sym, "direction": "BUY", "qty": abs(qty), "order_id": response["norenordno"], "type": "EXIT"})
                            else:
                                self._log(f"EXIT FAILED {sym}")
                        except Exception as e:
                            self._log(f"EXIT ERROR {sym}: {e}")

                self.on_positions(pos_list)
                self._log(f"Cycle #{cycle} done | Total PnL: {total_pnl}")

            except Exception as e:
                self._log(f"MONITOR ERROR: {e}")

            self._stop_event.wait(MONITOR_INTERVAL)
            cycle += 1